"""World Bank Commodity Price Data, "The Pink Sheet" (bonus B5): the monthly Excel file.

The official page https://www.worldbank.org/en/research/commodity-markets links the current
CMO-Historical-Data-Monthly.xlsx. The link moves to a new document about once a year (the old
one stays online, frozen), and the file is replaced early every month with the previous
month's averages. Its "Monthly Prices" sheet has a few title rows ("Updated on September 02,
2026"), a row of series names ("Rice, Thai 5% "), a row of units ("($/mt)"), then one row
per month ("2026M08") in nominal US dollars. Missing values are "…", ".." or empty.

Series are found by name and their unit is checked, so a moved column never shifts the
numbers and a changed unit never turns tons into kilograms."""

import io
import math
import re
import zipfile
from collections import Counter
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any
from urllib.parse import urljoin

import openpyxl

# The monthly file linked from the official page on 2026-09-20; used only when the page cannot
# be read and no address is known yet (the worker normally finds the current one itself).
KNOWN_MONTHLY_URL = (
    "https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026"
    "/related/CMO-Historical-Data-Monthly.xlsx"
)
MONTHLY_SHEET = "Monthly Prices"
MAX_UNPACKED = 64_000_000  # the real file unpacks to about 3.5 MB
HEADER_SEARCH_ROWS = 20

_LINK = re.compile(r"""href\s*=\s*["']([^"']*CMO-Historical-Data-Monthly\.xlsx)["']""", re.I)
_UPDATED = re.compile(r"Updated on\s+([A-Za-z]+ \d{1,2}, \d{4})")
_MONTH = re.compile(r"^(\d{4})M(\d{2})$")
_UNIT = re.compile(r"^\(\s*\$\s*/\s*([A-Za-z]+)\s*\)$")


class FormatError(ValueError):
    """The download is not the monthly workbook we know."""


@dataclass(frozen=True)
class SeriesSpec:
    """A series to read: our id, its name in the sheet and the unit its price must be in."""

    id: str
    column: str
    unit: str  # "mt" or "kg"


@dataclass(frozen=True)
class MonthlyPrice:
    series_id: str
    month: date  # first day of the month
    usd: float  # US dollars per the series' unit


@dataclass(frozen=True)
class PinkSheet:
    published: date | None  # "Updated on …" in the sheet
    prices: list[MonthlyPrice]
    cells: int  # month × series cells read
    dropped: dict[str, int] = field(default_factory=dict)  # reason → cells
    problems: dict[str, str] = field(default_factory=dict)  # series id → why it was left out


def find_monthly_url(html: str, page_url: str) -> str | None:
    """The monthly file's address as linked from the official page (absolute)."""
    match = _LINK.search(html)
    return urljoin(page_url, match.group(1).strip()) if match else None


def _name(value: Any) -> str:
    """Series name as a key: spaces collapsed, footnote marks ("Beef **") removed."""
    return " ".join(str(value).split()).rstrip("*").strip() if value is not None else ""


def _unit(value: Any) -> str | None:
    match = _UNIT.match(str(value).strip()) if value is not None else None
    return match.group(1).lower() if match else None


def _month(value: Any) -> date | None:
    match = _MONTH.match(str(value).strip()) if value is not None else None
    if not match or not 1 <= int(match.group(2)) <= 12:
        return None
    return date(int(match.group(1)), int(match.group(2)), 1)


def _published(rows: Iterable[Sequence[Any]]) -> date | None:
    for row in rows:
        for value in row:
            match = _UPDATED.search(str(value)) if isinstance(value, str) else None
            if match:
                try:
                    return datetime.strptime(match.group(1), "%B %d, %Y").date()
                except ValueError:
                    return None
    return None


def _check_size(content: bytes, max_unpacked: int) -> None:
    """Refuses anything but a zip (xlsx) of a sane unpacked size (no zip bombs)."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            unpacked = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise FormatError("not an Excel file") from exc
    if unpacked > max_unpacked:
        raise FormatError(f"file too large once unpacked ({unpacked} bytes)")


def _classify(value: Any) -> float | str:
    """A usable price, or the reason the cell is dropped."""
    if value is None or (isinstance(value, str) and value.strip() in {"", "…", "...", ".."}):
        return "missing"
    if isinstance(value, bool) or not isinstance(value, int | float):
        return "not_a_number"
    number = float(value)
    if not math.isfinite(number) or number <= 0:
        return "non_positive"
    return number


def parse_monthly(
    content: bytes, series: Sequence[SeriesSpec], *, max_unpacked: int = MAX_UNPACKED
) -> PinkSheet:
    """Every month of the given series. Series that cannot be read safely are left out and
    explained in `problems`; a workbook without any of them raises FormatError."""
    _check_size(content, max_unpacked)
    try:
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many kinds for a damaged file
        raise FormatError(f"unreadable workbook: {exc}") from exc
    try:
        if MONTHLY_SHEET not in book.sheetnames:
            raise FormatError(f"no sheet {MONTHLY_SHEET!r}")
        rows = [tuple(row) for row in book[MONTHLY_SHEET].iter_rows(values_only=True)]
    finally:
        book.close()

    wanted = {_name(s.column): s for s in series}
    header = next(
        (
            i
            for i, row in enumerate(rows[:HEADER_SEARCH_ROWS])
            if any(_name(v) in wanted for v in row)
        ),
        None,
    )
    if header is None or header + 1 >= len(rows):
        raise FormatError("none of the series in the sheet")
    names, units = rows[header], rows[header + 1]

    columns: dict[str, int] = {}
    problems: dict[str, str] = {}
    positions = {_name(v): i for i, v in enumerate(names) if _name(v)}
    for spec in series:
        index = positions.get(_name(spec.column))
        unit = _unit(units[index]) if index is not None and index < len(units) else None
        if index is None:
            problems[spec.id] = f"no column {spec.column!r}"
        elif unit != spec.unit:
            found = f"$/{unit}" if unit else "unknown"
            problems[spec.id] = f"unit is {found}, expected $/{spec.unit}"
        else:
            columns[spec.id] = index
    if not columns:
        raise FormatError(f"none of the series can be read: {problems}")

    prices: list[MonthlyPrice] = []
    dropped: Counter[str] = Counter()
    cells = 0
    months = 0
    for row in rows[header + 2 :]:
        month = _month(row[0]) if row else None
        if month is None:
            continue
        months += 1
        for series_id, index in columns.items():
            cells += 1
            value = _classify(row[index] if index < len(row) else None)
            if isinstance(value, str):
                dropped[value] += 1
            else:
                prices.append(MonthlyPrice(series_id, month, value))
    if months == 0:
        raise FormatError("no month rows in the sheet")
    return PinkSheet(
        published=_published(rows[:header]),
        prices=prices,
        cells=cells,
        dropped=dict(dropped),
        problems=problems,
    )
