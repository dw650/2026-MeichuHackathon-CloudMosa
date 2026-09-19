"""World Bank Pink Sheet (bonus B5): finding the monthly file and reading its series.

`tests/fixtures/wb_pink_monthly_sample.xlsx` is cut from the real CMO-Historical-Data-Monthly.xlsx
downloaded on 2026-09-20 ("Updated on September 02, 2026"): the "Mismatch Details" sheet's first
rows, then the "Monthly Prices" sheet with its six header rows, every column (72 series), and the
months 1960M01–M02 plus 2024M06–2026M08 as published. `wb_commodity_markets_excerpt.html` is the
"Pink Sheet Data" table of https://www.worldbank.org/en/research/commodity-markets that day."""

import io
from collections.abc import Callable
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.worksheet import Worksheet

from app.ingest.intl.pink_sheet import (
    FormatError,
    SeriesSpec,
    find_monthly_url,
    parse_monthly,
)
from app.seed.loader import load_intl_series

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"
SAMPLE = (FIXTURES / "wb_pink_monthly_sample.xlsx").read_bytes()
PAGE = (FIXTURES / "wb_commodity_markets_excerpt.html").read_text(encoding="utf-8")
PAGE_URL = "https://www.worldbank.org/en/research/commodity-markets"
CURRENT = (
    "https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026"
    "/related/CMO-Historical-Data-Monthly.xlsx"
)
SPECS = [SeriesSpec(s.id, s.column, s.unit) for s in load_intl_series().series]
AUG = date(2026, 8, 1)
# Values of August 2026 as published ($/mt; sugar $/kg).
AUGUST = {
    "rice": 471,
    "wheat": 330,
    "maize": 224,
    "soybeans": 482,
    "sugar": 0.38,
    "palm_oil": 1117,
}
HEADER_ROW, UNIT_ROW, FIRST_DATA_ROW = 5, 6, 7  # 1-based rows of the sheet


def variant(edit: Callable[[Worksheet], None]) -> bytes:
    """The sample with some cells of the monthly sheet changed."""
    book = openpyxl.load_workbook(io.BytesIO(SAMPLE))
    edit(book["Monthly Prices"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def column_of(sheet: Worksheet, name: str) -> int:
    for cell in sheet[HEADER_ROW]:
        if isinstance(cell.value, str) and cell.value.strip() == name:
            return int(cell.column)
    raise AssertionError(name)


def last_row(sheet: Worksheet) -> int:
    return int(sheet.max_row)


# ---------- the link on the official page ----------


def test_the_page_links_the_current_monthly_file() -> None:
    assert find_monthly_url(PAGE, PAGE_URL) == CURRENT


def test_a_relative_link_is_resolved_against_the_page() -> None:
    html = '<a href="/en/doc/x/related/CMO-Historical-Data-Monthly.xlsx">Monthly prices</a>'
    assert find_monthly_url(html, PAGE_URL) == (
        "https://www.worldbank.org/en/doc/x/related/CMO-Historical-Data-Monthly.xlsx"
    )


def test_a_page_without_the_link_gives_none() -> None:
    annual_only = PAGE.replace("CMO-Historical-Data-Monthly.xlsx", "CMO-Pink-Sheet.pdf")
    assert find_monthly_url(annual_only, PAGE_URL) is None


# ---------- the monthly sheet ----------


def test_the_sample_gives_every_month_of_the_six_series() -> None:
    sheet = parse_monthly(SAMPLE, SPECS)
    assert sheet.published == date(2026, 9, 2)
    assert sheet.problems == {}
    by_series: dict[str, dict[date, float]] = {}
    for price in sheet.prices:
        by_series.setdefault(price.series_id, {})[price.month] = price.usd
    assert set(by_series) == set(AUGUST)
    for series_id, value in AUGUST.items():
        months = by_series[series_id]
        assert months[AUG] == pytest.approx(value)
        assert len(months) == 29  # 1960M01, 1960M02 and 2024M06–2026M08
        assert min(months) == date(1960, 1, 1)
    assert by_series["rice"][date(2026, 7, 1)] == pytest.approx(467)
    assert by_series["wheat"][date(2025, 9, 1)] == pytest.approx(233.8)
    assert sheet.cells == 6 * 29
    assert sheet.dropped == {}


def test_columns_are_found_by_name_whatever_the_spaces_and_footnote_marks() -> None:
    def rename(ws: Worksheet) -> None:
        ws.cell(HEADER_ROW, column_of(ws, "Maize")).value = "  Maize  **"

    sheet = parse_monthly(variant(rename), SPECS)
    assert sheet.problems == {}
    assert {p.series_id for p in sheet.prices} == set(AUGUST)


def test_missing_and_bad_values_are_dropped_and_counted() -> None:
    def spoil(ws: Worksheet) -> None:
        rice, sugar, maize = (column_of(ws, n) for n in ("Rice, Thai 5%", "Sugar, world", "Maize"))
        end = last_row(ws)
        ws.cell(end, rice).value = "…"
        ws.cell(end - 1, rice).value = None
        ws.cell(end, sugar).value = 0
        ws.cell(end, maize).value = "n/a"

    sheet = parse_monthly(variant(spoil), SPECS)
    latest = {p.series_id for p in sheet.prices if p.month == AUG}
    assert latest == {"wheat", "soybeans", "palm_oil"}
    assert sheet.dropped == {"missing": 2, "non_positive": 1, "not_a_number": 1}
    assert sheet.cells == 6 * 29


def test_a_series_whose_unit_changed_is_left_out() -> None:
    def per_kg(ws: Worksheet) -> None:
        ws.cell(UNIT_ROW, column_of(ws, "Rice, Thai 5%")).value = "($/kg)"

    sheet = parse_monthly(variant(per_kg), SPECS)
    assert sheet.problems == {"rice": "unit is $/kg, expected $/mt"}
    assert "rice" not in {p.series_id for p in sheet.prices}
    assert len({p.series_id for p in sheet.prices}) == 5


def test_a_series_whose_column_is_gone_is_reported() -> None:
    def drop(ws: Worksheet) -> None:
        ws.cell(HEADER_ROW, column_of(ws, "Palm oil")).value = "Palm oil, crude"

    sheet = parse_monthly(variant(drop), SPECS)
    assert sheet.problems == {"palm_oil": "no column 'Palm oil'"}


def test_rows_that_are_not_months_are_skipped() -> None:
    def notes(ws: Worksheet) -> None:
        ws.cell(last_row(ws) + 2, 1).value = "Note: prices are nominal."
        ws.cell(FIRST_DATA_ROW, 1).value = "1960-01"  # not the YYYYMmm form

    sheet = parse_monthly(variant(notes), SPECS)
    assert date(1960, 1, 1) not in {p.month for p in sheet.prices}
    assert sheet.cells == 6 * 28


def test_the_update_date_is_optional() -> None:
    def undated(ws: Worksheet) -> None:
        ws.cell(4, 1).value = None

    assert parse_monthly(variant(undated), SPECS).published is None


@pytest.mark.parametrize(
    ("content", "message"),
    [
        (b"<html>maintenance</html>", "not an Excel file"),
        (SAMPLE, "file too large once unpacked"),
    ],
)
def test_files_that_cannot_be_read_are_refused(content: bytes, message: str) -> None:
    limit = 10_000 if content is SAMPLE else 50_000_000
    with pytest.raises(FormatError, match=message):
        parse_monthly(content, SPECS, max_unpacked=limit)


def test_a_workbook_without_the_monthly_sheet_is_refused() -> None:
    book = openpyxl.Workbook()
    book.active.title = "Monthly Indices"  # type: ignore[union-attr]
    out = io.BytesIO()
    book.save(out)
    with pytest.raises(FormatError, match="no sheet 'Monthly Prices'"):
        parse_monthly(out.getvalue(), SPECS)


def test_a_sheet_with_none_of_our_series_is_refused() -> None:
    def rename_all(ws: Worksheet) -> None:
        for cell in ws[HEADER_ROW]:
            if cell.value is not None:
                cell.value = f"{cell.value} (old)"

    with pytest.raises(FormatError, match="none of the series"):
        parse_monthly(variant(rename_all), SPECS)


def test_a_sheet_without_month_rows_is_refused() -> None:
    def no_months(ws: Worksheet) -> None:
        for row in range(FIRST_DATA_ROW, last_row(ws) + 1):
            ws.cell(row, 1).value = None

    with pytest.raises(FormatError, match="no month rows"):
        parse_monthly(variant(no_months), SPECS)
