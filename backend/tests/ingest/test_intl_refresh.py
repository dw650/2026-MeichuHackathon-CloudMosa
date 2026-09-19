"""Keeping the B5 data current: what is downloaded when, and what a failure leaves behind.
The three addresses are served by tests/intl_server.py from the saved real samples."""

import io
from collections.abc import AsyncIterator, Sequence
from datetime import UTC, date, datetime, timedelta
from typing import Any

import openpyxl
import pytest
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.config import Settings
from app.db.session import create_engine, create_sessionmaker
from app.ingest.intl.refresh import IntlConfig, refresh_intl
from app.ingest.pipeline import RunSummary
from app.seed.loader import load_intl_series
from app.seed.schema import IntlSeriesSeed
from tests.intl_server import (
    CONFIG,
    FILE_URL,
    FX_URL,
    PAGE,
    PAGE_URL,
    SAMPLE,
    IntlServer,
    fx_payload,
)

SEEDS = load_intl_series().series
T0 = datetime(2026, 9, 19, 6, 10, tzinfo=UTC)
AFTER_FX_UPDATE = datetime(2026, 9, 20, 0, 31, tzinfo=UTC)  # the sample announces 00:30:21
NEW_URL = FILE_URL.replace("74e8be41ceb20fa0da750cda2f6b9e4e-0050012026", "abc-0050012027")
OCTOBER = "Fri, 02 Oct 2026 20:00:00 GMT"


def unix(moment: datetime) -> int:
    return int(moment.timestamp())


class Sleeps(list[float]):
    async def __call__(self, seconds: float) -> None:
        self.append(seconds)


@pytest.fixture
async def maker(
    settings: Settings, clean_db: str
) -> AsyncIterator[async_sessionmaker[AsyncSession]]:
    engine = create_engine(settings)
    yield create_sessionmaker(engine)
    await engine.dispose()


async def refresh(
    maker: async_sessionmaker[AsyncSession],
    server: IntlServer,
    at: datetime,
    *,
    config: IntlConfig = CONFIG,
    seeds: Sequence[IntlSeriesSeed] = SEEDS,
    force: bool = False,
) -> list[tuple[str, str]]:
    summaries: list[RunSummary] = await refresh_intl(
        maker, seeds, config, lambda: at, transport=server.transport(), sleep=Sleeps(), force=force
    )
    return [(s.source, s.status) for s in summaries]


async def scalar(session: AsyncSession, sql: str, **params: Any) -> Any:
    return (await session.execute(text(sql), params)).scalar_one()


async def usd(session: AsyncSession, series: str, month: date) -> float:
    value = await scalar(
        session,
        "SELECT usd FROM intl_prices WHERE series_id = :s AND month = :m",
        s=series,
        m=month,
    )
    return float(value)


def workbook(edit: Any) -> bytes:
    """The sample with the monthly sheet changed by `edit(sheet)`."""
    book = openpyxl.load_workbook(io.BytesIO(SAMPLE))
    edit(book["Monthly Prices"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def column(sheet: Any, name: str) -> int:
    """1-based column of a series in the header row (row 5)."""
    return int(
        next(c.column for c in sheet[5] if isinstance(c.value, str) and c.value.strip() == name)
    )


async def test_the_first_run_downloads_each_source_once(
    maker: async_sessionmaker[AsyncSession], session: AsyncSession
) -> None:
    server = IntlServer()
    assert await refresh(maker, server, T0) == [("er_api", "ok"), ("wb_pink", "ok")]
    assert (server.hits(FX_URL), server.hits(PAGE_URL), server.hits(FILE_URL)) == (1, 1, 1)
    assert server.conditional(FILE_URL) == [False]

    assert await scalar(session, "SELECT count(*) FROM intl_series") == 6
    assert await scalar(session, "SELECT count(*) FROM intl_prices") == 6 * 29
    assert await usd(session, "rice", date(2026, 8, 1)) == 471
    assert await usd(session, "sugar", date(2026, 8, 1)) == 0.38
    rates = await session.execute(text("SELECT currency, per_usd, rate_date FROM fx_rates"))
    by_code = {c: (float(v), d) for c, v, d in rates.tuples().all()}
    assert by_code["TWD"] == (31.834145, date(2026, 9, 19))
    assert len(by_code) == 6

    sources = await session.execute(
        text(
            "SELECT id, url, last_modified, data_date, next_update_at, checked_at"
            " FROM intl_sources ORDER BY id"
        )
    )
    assert sources.tuples().all() == [
        (
            "er_api",
            FX_URL,
            None,
            date(2026, 9, 19),
            datetime(2026, 9, 20, 0, 30, 21, tzinfo=UTC),
            T0,
        ),
        ("wb_pink", FILE_URL, "Wed, 02 Sep 2026 20:17:37 GMT", date(2026, 9, 2), None, T0),
    ]
    run_rows = await session.execute(
        text("SELECT source, status, rows_in, rows_ok, rows_dropped FROM ingest_runs ORDER BY id")
    )
    assert run_rows.tuples().all() == [
        ("er_api", "ok", 6, 6, 0),
        ("wb_pink", "ok", 174, 174, 0),
    ]


async def test_restarts_download_nothing_until_a_source_is_due(
    maker: async_sessionmaker[AsyncSession], session: AsyncSession
) -> None:
    server = IntlServer()
    await refresh(maker, server, T0)
    assert await refresh(maker, server, T0 + timedelta(hours=1)) == []
    assert await refresh(maker, server, T0 + timedelta(hours=17)) == []
    assert len(server.requests) == 3

    # The provider's next rates are out: only the rates are fetched again.
    server.fx = fx_payload(
        time_last_update_unix=unix(datetime(2026, 9, 20, 0, 2, tzinfo=UTC)),
        time_next_update_unix=unix(datetime(2026, 9, 21, 0, 30, tzinfo=UTC)),
    )
    assert await refresh(maker, server, AFTER_FX_UPDATE) == [("er_api", "ok")]
    assert (server.hits(FX_URL), server.hits(PAGE_URL), server.hits(FILE_URL)) == (2, 1, 1)

    # A day later the Pink Sheet is checked: the page, then a conditional GET (304).
    day_later = T0 + timedelta(hours=23)
    assert await refresh(maker, server, day_later) == [("wb_pink", "ok")]
    assert (server.hits(PAGE_URL), server.conditional(FILE_URL)) == (2, [False, True])
    assert await scalar(session, "SELECT count(*) FROM intl_prices") == 6 * 29
    checked = "SELECT checked_at FROM intl_sources WHERE id = 'wb_pink'"
    assert await scalar(session, checked) == day_later
    last = "SELECT rows_in, rows_ok FROM ingest_runs ORDER BY id DESC LIMIT 1"
    assert (await session.execute(text(last))).tuples().one() == (0, 0)


async def test_rates_are_checked_a_day_later_even_without_an_announced_update(
    maker: async_sessionmaker[AsyncSession],
) -> None:
    server = IntlServer(fx=fx_payload(time_next_update_unix=None))
    await refresh(maker, server, T0)
    assert dict(await refresh(maker, server, T0 + timedelta(hours=23)))["wb_pink"] == "ok"
    assert await refresh(maker, server, T0 + timedelta(hours=24)) == [("er_api", "ok")]


async def test_a_late_provider_is_asked_at_most_once_an_hour(
    maker: async_sessionmaker[AsyncSession],
) -> None:
    server = IntlServer()  # always announces 2026-09-20 00:30:21, as if the update were late
    await refresh(maker, server, T0)
    assert await refresh(maker, server, AFTER_FX_UPDATE) == [("er_api", "ok")]
    assert await refresh(maker, server, AFTER_FX_UPDATE + timedelta(minutes=59)) == []
    after_an_hour = AFTER_FX_UPDATE + timedelta(hours=1)
    assert await refresh(maker, server, after_an_hour) == [("er_api", "ok")]


async def test_force_checks_both_sources(maker: async_sessionmaker[AsyncSession]) -> None:
    server = IntlServer()
    await refresh(maker, server, T0)
    done = await refresh(maker, server, T0 + timedelta(minutes=1), force=True)
    assert done == [("er_api", "ok"), ("wb_pink", "ok")]
    assert server.conditional(FILE_URL) == [False, True]


async def test_a_new_file_replaces_the_months(
    maker: async_sessionmaker[AsyncSession], session: AsyncSession
) -> None:
    server = IntlServer()
    await refresh(maker, server, T0)

    def revise(sheet: Any) -> None:
        sheet.cell(sheet.max_row, column(sheet, "Rice, Thai 5%")).value = 480
        sheet.cell(sheet.max_row - 1, column(sheet, "Maize")).value = "…"
        sheet.cell(4, 1).value = "Updated on October 02, 2026"

    server.files[FILE_URL] = workbook(revise)
    server.last_modified = OCTOBER
    assert dict(await refresh(maker, server, T0 + timedelta(hours=23)))["wb_pink"] == "ok"
    assert server.conditional(FILE_URL) == [False, True]
    assert await usd(session, "rice", date(2026, 8, 1)) == 480
    maize_july = "SELECT count(*) FROM intl_prices WHERE series_id = 'maize' AND month = :m"
    assert await scalar(session, maize_july, m=date(2026, 7, 1)) == 0
    state = "SELECT last_modified, data_date FROM intl_sources WHERE id = 'wb_pink'"
    assert (await session.execute(text(state))).tuples().one() == (OCTOBER, date(2026, 10, 2))


async def test_a_new_address_on_the_page_is_downloaded_in_full(
    maker: async_sessionmaker[AsyncSession], session: AsyncSession
) -> None:
    server = IntlServer()
    await refresh(maker, server, T0)
    server.page = PAGE.replace(FILE_URL, NEW_URL)
    server.files[NEW_URL] = SAMPLE
    assert dict(await refresh(maker, server, T0 + timedelta(hours=23)))["wb_pink"] == "ok"
    assert server.conditional(NEW_URL) == [False]
    assert await scalar(session, "SELECT url FROM intl_sources WHERE id = 'wb_pink'") == NEW_URL


async def test_when_the_page_fails_the_known_file_is_still_checked(
    maker: async_sessionmaker[AsyncSession],
) -> None:
    server = IntlServer()
    await refresh(maker, server, T0)
    server.down = {PAGE_URL}
    assert dict(await refresh(maker, server, T0 + timedelta(hours=23)))["wb_pink"] == "ok"
    assert server.hits(PAGE_URL) == 1 + 3  # retried, then given up
    assert server.conditional(FILE_URL) == [False, True]


async def test_without_the_page_the_built_in_address_is_used(
    maker: async_sessionmaker[AsyncSession], session: AsyncSession
) -> None:
    server = IntlServer(page="<html>maintenance</html>")
    assert await refresh(maker, server, T0) == [("er_api", "ok"), ("wb_pink", "ok")]
    assert server.hits(FILE_URL) == 1
    assert await scalar(session, "SELECT count(*) FROM intl_prices") == 6 * 29


async def test_a_configured_file_address_skips_the_page(
    maker: async_sessionmaker[AsyncSession],
) -> None:
    server = IntlServer(files={NEW_URL: SAMPLE})
    fixed = IntlConfig(page_url=PAGE_URL, file_url=NEW_URL, fallback_url=FILE_URL, fx_url=FX_URL)
    assert await refresh(maker, server, T0, config=fixed) == [("er_api", "ok"), ("wb_pink", "ok")]
    assert (server.hits(PAGE_URL), server.hits(NEW_URL), server.hits(FILE_URL)) == (0, 1, 0)


async def test_a_failed_download_keeps_the_data_and_is_tried_at_the_next_check(
    maker: async_sessionmaker[AsyncSession], session: AsyncSession
) -> None:
    server = IntlServer()
    await refresh(maker, server, T0)
    server.down = {FILE_URL}
    assert dict(await refresh(maker, server, T0 + timedelta(hours=23)))["wb_pink"] == "failed"
    assert await scalar(session, "SELECT count(*) FROM intl_prices") == 6 * 29
    assert await scalar(session, "SELECT checked_at FROM intl_sources WHERE id = 'wb_pink'") == T0
    failed = "SELECT error FROM ingest_runs WHERE status = 'failed'"
    assert "failed 3 attempts: HTTP 503" in await scalar(session, failed)

    server.down = set()
    assert dict(await refresh(maker, server, T0 + timedelta(hours=24)))["wb_pink"] == "ok"


async def test_a_file_we_cannot_read_changes_nothing(
    maker: async_sessionmaker[AsyncSession], session: AsyncSession
) -> None:
    server = IntlServer()
    await refresh(maker, server, T0)
    server.files[FILE_URL] = b"<html>Access denied</html>"
    server.last_modified = OCTOBER
    assert dict(await refresh(maker, server, T0 + timedelta(hours=23)))["wb_pink"] == "failed"
    state = "SELECT last_modified FROM intl_sources WHERE id = 'wb_pink'"
    assert await scalar(session, state) == "Wed, 02 Sep 2026 20:17:37 GMT"
    assert await usd(session, "rice", date(2026, 8, 1)) == 471


async def test_refused_rates_keep_the_old_ones(
    maker: async_sessionmaker[AsyncSession], session: AsyncSession
) -> None:
    server = IntlServer()
    await refresh(maker, server, T0)
    server.fx = {"result": "error", "error-type": "quota-reached"}
    assert await refresh(maker, server, AFTER_FX_UPDATE) == [("er_api", "failed")]
    twd = await scalar(session, "SELECT per_usd FROM fx_rates WHERE currency = 'TWD'")
    assert float(twd) == 31.834145


async def test_rates_that_are_not_json_fail_the_check(
    maker: async_sessionmaker[AsyncSession],
) -> None:
    class NotJson(IntlServer):
        def __call__(self, request: Any) -> Any:
            if str(request.url) == FX_URL:
                import httpx

                return httpx.Response(200, text="<html>oops</html>")
            return super().__call__(request)

    assert await refresh(maker, NotJson(), T0) == [("er_api", "failed"), ("wb_pink", "ok")]


async def test_a_series_missing_from_a_new_file_keeps_its_prices(
    maker: async_sessionmaker[AsyncSession], session: AsyncSession
) -> None:
    server = IntlServer()
    await refresh(maker, server, T0)

    def rename(sheet: Any) -> None:
        sheet.cell(5, column(sheet, "Palm oil")).value = "Palm oil, crude"

    server.files[FILE_URL] = workbook(rename)
    server.last_modified = OCTOBER
    assert dict(await refresh(maker, server, T0 + timedelta(hours=23)))["wb_pink"] == "ok"
    palm = "SELECT count(*) FROM intl_prices WHERE series_id = 'palm_oil'"
    assert await scalar(session, palm) == 29
    reasons = "SELECT drop_reasons FROM ingest_runs ORDER BY id DESC LIMIT 1"
    assert await scalar(session, reasons) == {"unreadable_series": 1}


async def test_the_series_follow_the_seed(
    maker: async_sessionmaker[AsyncSession], session: AsyncSession
) -> None:
    server = IntlServer()
    await refresh(maker, server, T0)
    fewer = [s for s in SEEDS if s.id != "sugar"]
    assert await refresh(maker, server, T0 + timedelta(hours=1), seeds=fewer) == []
    assert await scalar(session, "SELECT count(*) FROM intl_series") == 5
    assert await scalar(session, "SELECT count(*) FROM intl_prices") == 5 * 29
    order = await session.execute(text("SELECT id FROM intl_series ORDER BY sort"))
    assert order.scalars().all() == ["rice", "wheat", "maize", "soybeans", "palm_oil"]
